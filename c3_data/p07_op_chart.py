import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows
import os
import pandas as pd

cwd = os.getcwd()
filename = "financials.xlsx"
filepath = os.path.join(cwd, "data", filename)

df = pd.read_excel(filepath, header=0, index_col=0)
tdf = df.T

wb = openpyxl.Workbook()
ws = wb.create_sheet(index=0, title='Chart')
wb.remove(wb['Sheet'])

for row in dataframe_to_rows(tdf, index=True, header=True):
    if len(row) > 1:
        ws.append(row)
        print(row)
          
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "매출액/영업이익"
chart.y_axis.title = "금액(억원)"
chart.x_axis.title = "연도"

data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=5)
cats = Reference(ws, min_col=1, min_row=2, max_row=5)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.shape = 4
ws.add_chart(chart, "A8")

wb.save(os.path.join(os.getcwd(), "output", "financials_barchart.xlsx"))
print("\n*** output 폴더의 \"financials_barchart.xlsx\" 파일을 확인해 보세요. ***\n")
